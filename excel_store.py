from pathlib import Path

from openpyxl import Workbook, load_workbook

from storage_common import SHEET_TITLES, sheet_headers, sheet_values

DATA_DIR = Path(__file__).resolve().parent / "data"
EXCEL_PATH = DATA_DIR / "registrations.xlsx"


class ExcelStore:
    def upsert(self, record: dict) -> None:
        DATA_DIR.mkdir(exist_ok=True)
        event_type = record["event_type"]
        title = SHEET_TITLES[event_type]
        headers = sheet_headers(event_type)
        values = sheet_values(record)

        if EXCEL_PATH.exists():
            workbook = load_workbook(EXCEL_PATH)
        else:
            workbook = Workbook()
            first = workbook.active
            first.title = title

        if title in workbook.sheetnames:
            sheet = workbook[title]
        else:
            sheet = workbook.create_sheet(title)

        current_header = [cell.value for cell in sheet[1]]
        if current_header != headers:
            if sheet.max_row == 1 and all(value is None for value in current_header):
                pass
            else:
                sheet.delete_rows(1, 1)
                sheet.insert_rows(1)
            for index, column in enumerate(headers, start=1):
                sheet.cell(1, index, column)

        target = None
        for excel_row in range(2, sheet.max_row + 1):
            if str(sheet.cell(excel_row, 1).value) == str(record["user_id"]):
                target = excel_row
                break

        if target is None:
            sheet.append(values)
        else:
            for index, value in enumerate(values, start=1):
                sheet.cell(target, index, value)

        extra = [name for name in workbook.sheetnames if name not in SHEET_TITLES.values()]
        if extra and workbook.sheetnames[0] not in SHEET_TITLES.values():
            default = workbook.sheetnames[0]
            if workbook[default].max_row <= 1:
                workbook.remove(workbook[default])

        workbook.save(EXCEL_PATH)
